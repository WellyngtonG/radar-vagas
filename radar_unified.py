"""
Radar de Vagas — Descoberta + Ranqueamento por Afinidade
=========================================================
Varre LinkedIn, Indeed e Glassdoor com JobSpy, depois avalia
a afinidade de cada vaga com o perfil do usuário usando um
sistema de scoring baseado em keywords + heurísticas.

Uso:
    python radar.py                    # Busca completa
    python radar.py --termo "AI Eng"   # Busca específica
    python radar.py --apenas-remoto    # Só vagas remotas
    python radar.py --top 15           # Mostra top 15

Saída: relatório no terminal + arquivo Excel com todas as vagas ranqueadas.
"""

import argparse
import os
import re
import sys
import yaml
from datetime import datetime, timedelta
from pathlib import Path

try:
    from jobspy import scrape_jobs
except ImportError:
    print("[ERRO] python-jobspy não instalado. Execute:")
    print("  uv pip install python-jobspy")
    sys.exit(1)

try:
    import openpyxl
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False


# ── Carregar Perfil ──────────────────────────────────────────

import hashlib


def _gerar_id_vaga(vaga: dict) -> str:
    """Gera um hash único de 12 caracteres para a vaga (URL ou Título+Empresa)."""
    url = str(vaga.get("job_url", "")).strip()
    if url:
        return hashlib.md5(url.encode("utf-8")).hexdigest()[:12]
    raw = f"{vaga.get('title', '')}_{vaga.get('company', '')}".strip()
    return hashlib.md5(raw.encode("utf-8")).hexdigest()[:12]


def carregar_historico() -> set:
    """Carrega IDs de vagas já processadas do arquivo historico_vagas.txt."""
    hist_file = Path(__file__).parent / "historico_vagas.txt"
    if hist_file.exists():
        with open(hist_file, "r", encoding="utf-8") as f:
            return {line.strip() for line in f if line.strip()}
    return set()


def salvar_historico(novos_ids: set):
    """Salva novos IDs de vagas no arquivo historico_vagas.txt."""
    if not novos_ids:
        return
    hist_file = Path(__file__).parent / "historico_vagas.txt"
    with open(hist_file, "a", encoding="utf-8") as f:
        for vid in sorted(novos_ids):
            f.write(f"{vid}\n")


def carregar_perfil(path: str = "perfil.yaml") -> dict:
    """Carrega o perfil profissional do arquivo YAML e normaliza o schema."""
    perfil_path = Path(__file__).parent / path
    if not perfil_path.exists():
        print(f"[ERRO] Arquivo de perfil não encontrado: {perfil_path}")
        print("Crie o arquivo perfil.yaml com suas especificações.")
        sys.exit(1)

    with open(perfil_path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}

    skills = data.get("skills", [])
    skills_basico = data.get("skills_basico", [])
    habs = data.get("habilidades", {})
    if isinstance(habs, dict):
        extra_skills = []
        for cat, itens in habs.items():
            if isinstance(itens, list):
                extra_skills.extend(itens)
        if not skills:
            skills = extra_skills
        else:
            for s in extra_skills:
                if s not in skills:
                    skills.append(s)
    elif isinstance(habs, list) and not skills:
        skills = habs
    data["skills"] = skills
    data["skills_basico"] = skills_basico

    if "termos_busca" in data and "areas_interesse" not in data:
        data["areas_interesse"] = data["termos_busca"]

    if "sites_busca" in data and "sites" not in data:
        data["sites"] = data["sites_busca"]

    filtros = data.get("filtros", {})
    if isinstance(filtros, dict):
        if "max_resultados_por_termo" in filtros and "max_resultados_por_busca" not in data:
            data["max_resultados_por_busca"] = filtros["max_resultados_por_termo"]
        if "excluir_termos" in filtros and "excluir_termos" not in data:
            data["excluir_termos"] = filtros["excluir_termos"]

    return data


# ── Integração 99freelas ──────────────────────────────────
def scrapear_99freelas(categoria="web-mobile-e-software") -> list:
    """Scrapeia projetos do 99freelas usando requests/BeautifulSoup."""
    import requests
    from bs4 import BeautifulSoup
    from datetime import datetime

    url = f"https://www.99freelas.com.br/projects?categoria={categoria}"
    headers = {
        "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }

    print(f"=== Scraping 99freelas: {url} ===")
    try:
        resp = requests.get(url, headers=headers, timeout=15)
        print(f"Status HTTP: {resp.status_code} | Tamanho: {len(resp.text)} bytes")
        if resp.status_code != 200:
            return []
        html_content = resp.text
    except Exception as e:
        print(f"Erro na requisicao: {e}")
        return []

    soup = BeautifulSoup(html_content, "html.parser")
    items = soup.select(".result-item")
    print(f"Projetos encontrados: {len(items)}")

    projetos = []
    for i, item in enumerate(items, 1):
        pid = item.get("data-id", "")
        pname = item.get("data-nome", "")
        title_a = item.select_one("h1.title a")
        title = title_a.text.strip() if title_a else pname
        link = ("https://www.99freelas.com.br" + title_a["href"]) if title_a and title_a.has_attr("href") else ""
        desc_elem = item.select_one(".description")
        desc = desc_elem.text.strip() if desc_elem else ""
        skills = [h.text.strip() for h in item.select(".habilidade")]
        info_elem = item.select_one(".information")
        info = " ".join(info_elem.text.split()) if info_elem else ""
        dt_elem = item.select_one(".datetime")
        dt_str = ""
        if dt_elem and dt_elem.has_attr("cp-datetime"):
            try:
                ts = int(dt_elem["cp-datetime"]) / 1000
                dt_str = datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M")
            except Exception:
                pass

        projetos.append({
            "title": title,
            "description": desc,
            "job_url": link,
            "site": "99freelas",
            "skills": skills,
            "meta": info,
            "date_posted": dt_str,
        })

    return projetos

# ── Motor de Match (Scoring) ────────────────────────────────

def normalizar(texto) -> str:
    """Normaliza texto para comparação case-insensitive."""
    if texto is None or (isinstance(texto, float) and str(texto) == "nan"):
        return ""
    texto = str(texto)
    if not texto:
        return ""
    return texto.lower().strip()


def _match_variante(variante: str, texto: str) -> bool:
    """Verifica match no texto usando boundary de palavra para termos curtos (<= 4 chars)."""
    v = normalizar(variante)
    if not v:
        return False
    if len(v) <= 4 or not re.search(r'\w', v):
        return bool(re.search(r'\b' + re.escape(v) + r'\b', texto))
    return v in texto


def calcular_match(vaga: dict, perfil: dict) -> dict:
    """
    Calcula a afinidade (0-100) entre uma vaga e o perfil.
    Retorna dict com score, detalhes (hits/misses), e justificativa.
    """
    titulo = normalizar(vaga.get("title", ""))
    descricao = normalizar(vaga.get("description", ""))
    texto_completo = f"{titulo} {descricao}"
    
    score = 0
    max_score = 0
    hits = []      # Skills encontradas na vaga
    misses = []    # Skills da vaga que o candidato não tem
    bonus = []     # Pontos extras

    # ── 1. Skills Principais e Básicas (com Teto Normalizado) ──
    skills_principais = [normalizar(s) for s in perfil.get("skills", [])]
    skills_score_raw = 0
    for skill in skills_principais:
        variantes = _gerar_variantes(skill)
        if any(_match_variante(v, texto_completo) for v in variantes):
            skills_score_raw += 5
            hits.append(skill)

    # ── 2. Skills Básicas (peso 2 cada) ───────
    skills_basicas = [normalizar(s) for s in perfil.get("skills_basico", [])]
    for skill in skills_basicas:
        variantes = _gerar_variantes(skill)
        if any(_match_variante(v, texto_completo) for v in variantes):
            skills_score_raw += 2
            hits.append(f"{skill} (básico)")

    # Teto de pontuação de skills: 50 pontos (atingível com ~8-10 skills requeridas)
    teto_skills = 50
    score += min(skills_score_raw, teto_skills)
    max_score += teto_skills

    # ── 3. Senioridade (até 15 pontos) ───────────────────────
    raw_sen = perfil.get("senioridade", [])
    if isinstance(raw_sen, str):
        raw_sen = [s.strip() for s in re.split(r'[/,]', raw_sen)]
    elif not isinstance(raw_sen, list):
        raw_sen = [str(raw_sen)]

    sen_alvo = set()
    for s in raw_sen:
        sn = normalizar(s)
        if any(x in sn for x in ["junior", "júnior", "jr", "trainee"]):
            sen_alvo.add("junior")
        if any(x in sn for x in ["pleno", "mid"]):
            sen_alvo.add("pleno")
        if any(x in sn for x in ["senior", "sênior", "sr", "lead"]):
            sen_alvo.add("senior")

    max_score += 15
    mapa_senioridade = {
        "junior": ["junior", "júnior", "jr", "entry", "trainee"],
        "pleno": ["pleno", "mid", "mid-level", "midlevel", "intermediário", "intermediate"],
        "senior": ["senior", "sênior", "sr", "lead", "principal", "staff"],
    }
    nivel_vaga = "indefinido"
    for nivel, termos in mapa_senioridade.items():
        if any(_match_variante(t, titulo) for t in termos) or any(_match_variante(t, descricao) for t in termos):
            nivel_vaga = nivel
            break

    if nivel_vaga == "indefinido":
        score += 10  # Não especificado = neutro/bom
        bonus.append("Nível não especificado (neutro)")
    elif nivel_vaga in sen_alvo:
        score += 15
        bonus.append(f"Nível compatível ({nivel_vaga})")
    elif "pleno" in sen_alvo and nivel_vaga == "junior":
        score += 12  # Overqualified mas candidatável
        bonus.append("Overqualified (júnior, você é pleno)")
    elif "pleno" in sen_alvo and nivel_vaga == "senior":
        score += 5   # Stretch mas possível
        bonus.append("Stretch (sênior, pode tentar)")
    else:
        score += 3

    # ── 4. Modelo de Trabalho (até 10 pontos) ────────────────
    raw_mod = perfil.get("modelo_trabalho", [])
    if isinstance(raw_mod, str):
        raw_mod = [m.strip() for m in re.split(r'[/,]', raw_mod)]
    elif not isinstance(raw_mod, list):
        raw_mod = [str(raw_mod)]
    modos_alvo = [normalizar(m) for m in raw_mod]
    aceita_remoto = any("remot" in m for m in modos_alvo)
    aceita_hibrido = any("hibrid" in m or "híbrid" in m for m in modos_alvo)
    aceita_presencial = any("presenc" in m for m in modos_alvo)

    max_score += 10
    is_remote = vaga.get("is_remote", False)
    eh_remoto = bool(is_remote or "remote" in texto_completo or "remoto" in texto_completo)
    eh_hibrido = bool("hybrid" in texto_completo or "híbrido" in texto_completo or "hibrido" in texto_completo)

    if eh_remoto and aceita_remoto:
        score += 10
        bonus.append("✅ Remoto confirmado")
    elif eh_hibrido and aceita_hibrido:
        score += 8
        bonus.append("🏢 Híbrido aceito")
    elif not eh_remoto and not eh_hibrido:
        if aceita_presencial:
            score += 7
            bonus.append("📍 Presencial aceito")
        else:
            score += 3
            bonus.append("❓ Modelo não especificado")
    else:
        score += 3

    # ── 5. Termos de Exclusão (penalidade) ───────────────────
    excluir = [normalizar(t) for t in perfil.get("excluir_termos", [])]
    for termo in excluir:
        if not termo:
            continue
        if _match_variante(termo, titulo):
            score -= 30
            misses.append(f"⛔ Excluído no título: {termo}")
        elif _match_variante(termo, descricao):
            score -= 15
            misses.append(f"⛔ Excluído na descrição: {termo}")

    # ── 6. Bônus por Indicadores de Qualidade ────────────────
    if any(t in texto_completo for t in ["open source", "opensource", "código aberto"]):
        score += 3
        bonus.append("Open source mencionado")
    if any(t in texto_completo for t in ["startup", "scale-up", "scaleup"]):
        score += 2
        bonus.append("Startup/Scale-up")
    if "equity" in texto_completo or "stock" in texto_completo:
        score += 2
        bonus.append("Equity/Stock mencionado")

    # ── Normalização Final ───────────────────────────────────
    # Score como % do máximo atingível
    if max_score > 0:
        score_pct = max(0, min(100, int((score / max_score) * 100)))
    else:
        score_pct = 0

    # Extrair requisitos que o candidato talvez não tenha
    termos_exigidos_comuns = [
        "kubernetes", "k8s", "terraform", "java", "c#", "c++", "go ", "golang",
        "rust", "scala", "ruby", "php", "angular", "vue", "swift", "kotlin",
        "gcp", "azure", "machine learning", "deep learning", "data science",
        "spark", "hadoop", "kafka", "elasticsearch"
    ]
    todas_skills = set(skills_principais + skills_basicas)
    for req in termos_exigidos_comuns:
        if req in texto_completo and req not in todas_skills:
            misses.append(req)

    return {
        "score": score_pct,
        "hits": hits,
        "misses": misses[:5],  # Limitar para não poluir
        "bonus": bonus,
        "nivel_vaga": nivel_vaga,
    }


def _gerar_variantes(skill: str) -> list:
    """Gera variantes de busca para uma skill (case insensitive)."""
    variantes = [skill]
    # node.js → nodejs, node
    if "." in skill:
        variantes.append(skill.replace(".", ""))
        variantes.append(skill.split(".")[0])
    # ci/cd → cicd, ci cd
    if "/" in skill:
        variantes.append(skill.replace("/", ""))
        variantes.append(skill.replace("/", " "))
    # llms / ia generativa → llm, ia, generativa
    if " / " in skill:
        partes = skill.split(" / ")
        variantes.extend(partes)
    # Siglas comuns
    mapa = {
        "javascript": ["js", "javascript", "ecmascript"],
        "typescript": ["ts", "typescript"],
        "python": ["python", "py"],
        "docker": ["docker", "container", "containerização"],
        "automação": ["automação", "automation", "automatização"],
        "web scraping": ["scraping", "web scraping", "crawler", "crawling"],
        "apis rest": ["api", "rest", "restful", "api rest"],
        "bots (telegram, discord)": ["bot", "telegram", "discord", "chatbot"],
        "llms / ia generativa": ["llm", "llms", "ia generativa", "generative ai",
                                  "ai", "artificial intelligence", "gpt", "gemini",
                                  "claude", "openai", "langchain"],
        "prompt engineering": ["prompt", "prompt engineering", "prompting"],
        "git": ["git", "version control", "controle de versão"],
        "github": ["github", "gh"],
        "linux": ["linux", "ubuntu", "debian", "centos"],
        "windows": ["windows"],
        "selenium": ["selenium", "webdriver"],
        "playwright": ["playwright"],
        "ci/cd": ["ci/cd", "cicd", "ci cd", "continuous integration",
                   "continuous delivery", "github actions", "jenkins", "pipeline"],
        "react": ["react", "reactjs", "react.js"],
        "next.js": ["next.js", "nextjs", "next"],
        "sql": ["sql", "mysql", "sqlite"],
        "postgresql": ["postgresql", "postgres", "psql"],
        "mongodb": ["mongodb", "mongo", "nosql"],
        "aws": ["aws", "amazon web services", "s3", "ec2", "lambda"],
        "power bi": ["power bi", "powerbi", "pbi"],
        "power query": ["power query", "powerquery", "pq"],
        "dax": ["dax"],
        "excel avançado": ["excel", "excel avançado", "planilhas", "vba", "macros"],
        "tabelas dinâmicas": ["tabela dinâmica", "tabelas dinâmicas", "pivot table", "pivot tables"],
        "procv / procx": ["procv", "procx", "xlookup", "vlookup"],
        "análise de dados": ["análise de dados", "analise de dados", "data analysis", "analytics"],
        "business intelligence (bi)": ["bi", "business intelligence", "inteligência de negócios"],
        "regras de negócio e kpis": ["kpi", "kpis", "regras de negócio", "indicadores", "métricas"],
        "automação com ia": ["automação", "ia", "ia generativa", "automação com ia"],
        "custos e indústria": ["custos", "custo", "indústria", "industrial", "precificação"],
        "logística e operações": ["logística", "logistica", "operações", "operacoes", "supply chain"],
    }
    if skill in mapa:
        variantes.extend(mapa[skill])
    return list(set(variantes))


# ── Buscar Vagas ─────────────────────────────────────────────

def buscar_vagas(perfil: dict, termo_extra: str = None) -> list:
    """Busca vagas em múltiplos sites usando JobSpy + 99freelas."""
    sites = perfil.get("sites", ["indeed", "linkedin"])
    max_results = perfil.get("max_resultados_por_busca", 25)
    localizacao = perfil.get("localizacao", "Brazil")
    buscar_intl = perfil.get("buscar_internacional", False)

    # Termos de busca
    termos = perfil.get("areas_interesse", ["Python Developer"])
    if termo_extra:
        termos = [termo_extra]

    todas_vagas = []
    vagas_vistas = set()  # Deduplica por URL

    # Busca via JobSpy
    for termo in termos:
        # Busca BR
        print(f"\n🔍 Buscando: \"{termo}\" em {', '.join(sites)} ({localizacao})...")
        try:
            pais_indeed = perfil.get("pais_indeed")
            if not pais_indeed:
                paises_str = f"{localizacao} {perfil.get('paises_interesse', '')}".lower()
                pais_indeed = "brazil" if any(x in paises_str for x in ["brasil", "brazil"]) else "brazil"

            mod_trab = perfil.get("modelo_trabalho", "")
            if isinstance(mod_trab, list):
                so_remoto = len(mod_trab) == 1 and str(mod_trab[0]).strip().lower() == "remoto"
            else:
                so_remoto = str(mod_trab).strip().lower() == "remoto"

            jobs = scrape_jobs(
                site_name=sites,
                search_term=termo,
                location=localizacao,
                results_wanted=max_results,
                country_indeed=str(pais_indeed).lower(),
                is_remote=so_remoto,
            )
            for _, row in jobs.iterrows():
                url = str(row.get("job_url", ""))
                if url and url not in vagas_vistas:
                    vagas_vistas.add(url)
                    todas_vagas.append(row.to_dict())
            print(f"   → {len(jobs)} vagas encontradas para \"{termo}\"")
        except Exception as e:
            print(f"   ⚠️ Erro ao buscar \"{termo}\": {e}")

        # Busca internacional (se habilitada)
        if buscar_intl and termo_extra is None:
            termo_en = _traduzir_termo(termo)
            if termo_en != termo:
                print(f"🌍 Buscando internacional: \"{termo_en}\"...")
                try:
                    jobs_intl = scrape_jobs(
                        site_name=["indeed", "linkedin"],
                        search_term=termo_en,
                        location="Remote",
                        results_wanted=max_results,
                        is_remote=True,
                    )
                    for _, row in jobs_intl.iterrows():
                        url = str(row.get("job_url", ""))
                        if url and url not in vagas_vistas:
                            vagas_vistas.add(url)
                            todas_vagas.append(row.to_dict())
                    print(f"   → {len(jobs_intl)} vagas internacionais")
                except Exception as e:
                    print(f"   ⚠️ Erro na busca internacional: {e}")

    # Busca 99freelas (projetos de freelance)
    print("\n🔍 Buscando projetos 99freelas...")
    cat_99 = perfil.get("categoria_99freelas", "web-mobile-e-software")
    if not cat_99 or str(cat_99).lower() in ["none", "null", "false", "nao", "não", ""]:
        print("   ℹ️ Busca 99freelas desativada no perfil.")
        projetos_99freelas = []
    else:
        print(f"   (categoria: {cat_99})")
        projetos_99freelas = scrapear_99freelas(categoria=str(cat_99).strip())
    for proj in projetos_99freelas:
        url = proj.get("job_url", "")
        if url and url not in vagas_vistas:
            vagas_vistas.add(url)
            # Converter para formato compatível com o restodo do radar
            proj_formatado = {
                "title": proj.get("title", ""),
                "description": proj.get("description", ""),
                "job_url": url,
                "site": "99freelas",
                "is_remote": True,  # 99freelas é remoto por padrão
                "company": "Freelancer",
                "location": localizacao,
            }
            todas_vagas.append(proj_formatado)
    print(f"   → {len(projetos_99freelas)} projetos 99freelas encontrados")

    return todas_vagas


def _traduzir_termo(termo: str) -> str:
    """Tradução simples de termos de busca PT→EN."""
    mapa = {
        "desenvolvedor python": "Python Developer",
        "automação": "Automation Engineer",
        "backend developer": "Backend Developer",
    }
    return mapa.get(termo.lower(), termo)


# ── Ranquear e Formatar ─────────────────────────────────────

def ranquear_vagas(vagas: list, perfil: dict) -> list:
    """Calcula match score e ordena por afinidade."""
    resultados = []
    for vaga in vagas:
        match = calcular_match(vaga, perfil)
        resultados.append({
            "score": match["score"],
            "titulo": vaga.get("title", "N/A"),
            "empresa": vaga.get("company", "N/A"),
            "local": vaga.get("location", "N/A"),
            "remoto": "✅" if vaga.get("is_remote") else "❌",
            "salario": _formatar_salario(vaga),
            "data": str(vaga.get("date_posted", "N/A")),
            "url": vaga.get("job_url", ""),
            "site": vaga.get("site", ""),
            "hits": ", ".join(match["hits"][:6]),
            "misses": ", ".join(match["misses"][:3]),
            "bonus": " | ".join(match["bonus"]),
            "nivel": match["nivel_vaga"],
            "descricao_curta": _resumir_descricao(vaga.get("description", "")),
        })

    resultados.sort(key=lambda x: x["score"], reverse=True)
    return resultados


def _formatar_salario(vaga: dict) -> str:
    """Formata faixa salarial se disponível."""
    min_val = vaga.get("min_amount")
    max_val = vaga.get("max_amount")
    moeda = vaga.get("currency", "")
    intervalo = vaga.get("interval", "")

    def _valido(v):
        return isinstance(v, (int, float)) and v == v and v > 0

    if _valido(min_val) and _valido(max_val):
        return f"{moeda} {min_val:,.0f}-{max_val:,.0f}/{intervalo}"
    elif _valido(min_val):
        return f"{moeda} {min_val:,.0f}+/{intervalo}"
    return "Não informado"


def _resumir_descricao(desc: str, max_chars: int = 200) -> str:
    """Extrai as primeiras linhas relevantes da descrição."""
    if not desc or not isinstance(desc, str):
        return ""
    # Remove tags HTML básicas se houver
    desc = re.sub(r'<[^>]+>', ' ', desc)
    desc = re.sub(r'\s+', ' ', desc).strip()
    if len(desc) > max_chars:
        return desc[:max_chars] + "..."
    return desc


# ── Exportar Resultados ─────────────────────────────────────

def exibir_terminal(resultados: list, top_n: int = 20):
    """Exibe os resultados ranqueados no terminal."""
    print("\n" + "=" * 70)
    print(f"  🎯 RADAR DE VAGAS — Top {min(top_n, len(resultados))} de {len(resultados)} vagas")
    print("=" * 70)

    for i, r in enumerate(resultados[:top_n], 1):
        # Cor do score
        if r["score"] >= 75:
            indicador = "🟢"
        elif r["score"] >= 50:
            indicador = "🟡"
        elif r["score"] >= 30:
            indicador = "🟠"
        else:
            indicador = "🔴"

        print(f"\n{indicador} #{i} — Match: {r['score']}%")
        print(f"   📌 {r['titulo']}")
        print(f"   🏢 {r['empresa']} | 📍 {r['local']} | Remoto: {r['remoto']}")
        print(f"   💰 {r['salario']} | 📅 {r['data']}")
        print(f"   🔗 {r['url']}")
        if r["hits"]:
            print(f"   ✅ Skills match: {r['hits']}")
        if r["misses"]:
            print(f"   ⚠️  Gaps: {r['misses']}")
        if r["bonus"]:
            print(f"   💡 {r['bonus']}")
        print(f"   📝 {r['descricao_curta']}")

    # Estatísticas
    print("\n" + "-" * 70)
    scores = [r["score"] for r in resultados]
    verde = sum(1 for s in scores if s >= 75)
    amarelo = sum(1 for s in scores if 50 <= s < 75)
    laranja = sum(1 for s in scores if 30 <= s < 50)
    vermelho = sum(1 for s in scores if s < 30)
    print(f"📊 Estatísticas: 🟢 {verde} ótimas | 🟡 {amarelo} boas | 🟠 {laranja} parciais | 🔴 {vermelho} fracas")
    print("-" * 70)


def exportar_excel(resultados: list, output_dir: str = "."):
    """Exporta resultados para Excel com formatação."""
    if not HAS_EXCEL:
        print("⚠️ openpyxl não instalado. Exportando como CSV...")
        return exportar_csv(resultados, output_dir)

    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filepath = Path(output_dir) / f"radar_vagas_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vagas Ranqueadas"

    # Headers
    headers = ["#", "Score", "Título", "Empresa", "Local", "Remoto",
               "Salário", "Data", "Skills Match", "Gaps", "Nível", "URL"]
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Dados
    for i, r in enumerate(resultados, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        score_cell = ws.cell(row=row, column=2, value=r["score"])

        # Cor condicional do score
        if r["score"] >= 75:
            score_cell.fill = PatternFill(start_color="27ae60", end_color="27ae60", fill_type="solid")
            score_cell.font = Font(color="FFFFFF", bold=True)
        elif r["score"] >= 50:
            score_cell.fill = PatternFill(start_color="f39c12", end_color="f39c12", fill_type="solid")
            score_cell.font = Font(bold=True)
        elif r["score"] >= 30:
            score_cell.fill = PatternFill(start_color="e67e22", end_color="e67e22", fill_type="solid")
        else:
            score_cell.fill = PatternFill(start_color="e74c3c", end_color="e74c3c", fill_type="solid")
            score_cell.font = Font(color="FFFFFF")

        ws.cell(row=row, column=3, value=r["titulo"])
        ws.cell(row=row, column=4, value=r["empresa"])
        ws.cell(row=row, column=5, value=r["local"])
        ws.cell(row=row, column=6, value=r["remoto"])
        ws.cell(row=row, column=7, value=r["salario"])
        ws.cell(row=row, column=8, value=r["data"])
        ws.cell(row=row, column=9, value=r["hits"])
        ws.cell(row=row, column=10, value=r["misses"])
        ws.cell(row=row, column=11, value=r["nivel"])
        ws.cell(row=row, column=12, value=r["url"])

    # Auto-width
    for col in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col).value or ""))
            for r in range(1, min(len(resultados) + 2, 50))
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 50)

    wb.save(filepath)
    print(f"\n📁 Excel exportado: {filepath}")
    return str(filepath)


def exportar_csv(resultados: list, output_dir: str = ".") -> str:
    """Fallback: exporta como CSV."""
    import csv
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")
    filepath = Path(output_dir) / f"radar_vagas_{timestamp}.csv"

    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "score", "titulo", "empresa", "local", "remoto",
            "salario", "data", "hits", "misses", "nivel", "url"
        ])
        writer.writeheader()
        writer.writerows([{k: r[k] for k in writer.fieldnames} for r in resultados])

    print(f"\n📁 CSV exportado: {filepath}")
    return str(filepath)


# ── Gerar Resumo para Telegram ──────────────────────────────

def gerar_resumo_telegram(resultados: list, top_n: int = 10, top_freelas: int = 3) -> str:
    """Gera um resumo formatado para envio no Telegram separando vagas e freelas."""
    linhas = [f"🎯 **Radar de Vagas** — {datetime.now().strftime('%d/%m %H:%M')}"]

    vagas_gerais = [r for r in resultados if r.get("site") != "99freelas"]
    freelas = [r for r in resultados if r.get("site") == "99freelas"]

    linhas.append(f"Total: {len(resultados)} oportunidades ({len(vagas_gerais)} vagas | {len(freelas)} freelas)\n")
    linhas.append(f"📌 **Top {min(top_n, len(vagas_gerais))} Vagas:**\n")

    for i, r in enumerate(vagas_gerais[:top_n], 1):
        if r["score"] >= 75:
            icone = "🟢"
        elif r["score"] >= 50:
            icone = "🟡"
        else:
            icone = "🟠"

        linhas.append(f"{icone} **{r['score']}%** — {r['titulo']}")
        linhas.append(f"   🏢 {r['empresa']} | 📍 {r['local']}")
        if r["hits"]:
            linhas.append(f"   ✅ {r['hits']}")
        linhas.append(f"   🔗 {r['url']}\n")

    if freelas:
        linhas.append(f"💼 **Top {min(top_freelas, len(freelas))} Freelas (99freelas):**\n")
        for i, r in enumerate(freelas[:top_freelas], 1):
            if r["score"] >= 50:
                icone = "🟢"
            elif r["score"] >= 25:
                icone = "🟡"
            else:
                icone = "🟠"

            linhas.append(f"{icone} **{r['score']}%** — {r['titulo']}")
            if r["hits"]:
                linhas.append(f"   ✅ {r['hits']}")
            linhas.append(f"   🔗 {r['url']}\n")

    # Stats
    verde = sum(1 for r in resultados if r["score"] >= 75)
    amarelo = sum(1 for r in resultados if 50 <= r["score"] < 75)
    linhas.append(f"📊 {verde} ótimas | {amarelo} boas | {len(resultados) - verde - amarelo} parciais")

    return "\n".join(linhas)


# ── Main ─────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="🎯 Radar de Vagas — Busca e ranqueamento por afinidade")
    parser.add_argument("--termo", type=str, help="Termo de busca específico (ignora perfil)")
    parser.add_argument("--apenas-remoto", action="store_true", help="Filtrar apenas vagas remotas")
    parser.add_argument("--top", type=int, default=20, help="Quantidade de vagas no relatório (default: 20)")
    parser.add_argument("--sem-excel", action="store_true", help="Não exportar Excel")
    parser.add_argument("--telegram", action="store_true", help="Gerar resumo para Telegram")
    args = parser.parse_args()

    print("🎯 Radar de Vagas — Iniciando...")
    print("-" * 40)

    # 1. Carregar perfil
    perfil = carregar_perfil()
    print(f"👤 Perfil: {perfil.get('nome', 'N/A')}")
    print(f"🛠️  Skills: {', '.join(perfil.get('skills', [])[:8])}...")
    print(f"🎯 Nível: {perfil.get('senioridade', 'N/A')} | Modelo: {perfil.get('modelo_trabalho', 'N/A')}")

    # Override de remoto
    if args.apenas_remoto:
        perfil["modelo_trabalho"] = "remoto"

    # 2. Buscar vagas
    vagas = buscar_vagas(perfil, termo_extra=args.termo)

    if not vagas:
        print("\n❌ Nenhuma vaga encontrada. Tente ajustar os termos de busca no perfil.yaml")
        sys.exit(0)

    print(f"\n📦 Total de vagas coletadas (deduplicadas): {len(vagas)}")

    # 3. Ranquear
    print("⚙️  Calculando afinidade...")
    resultados = ranquear_vagas(vagas, perfil)

    # 4. Exibir
    exibir_terminal(resultados, top_n=args.top)

    # 5. Exportar
    if not args.sem_excel:
        output_dir = Path(__file__).parent
        exportar_excel(resultados, str(output_dir))

    # 6. Telegram
    if args.telegram:
        resumo = gerar_resumo_telegram(resultados, top_n=10)
        telegram_file = Path(__file__).parent / "ultimo_resumo_telegram.txt"
        with open(telegram_file, "w", encoding="utf-8") as f:
            f.write(resumo)
        print(f"\n📱 Resumo Telegram salvo em: {telegram_file}")
        print("\n--- PREVIEW TELEGRAM ---")
        print(resumo)

    print(f"\n✅ Concluído! {len(resultados)} vagas analisadas.")


if __name__ == "__main__":
    main()
